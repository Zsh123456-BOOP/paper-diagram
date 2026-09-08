"""Contract tests for the reused native PowerPoint renderer."""

from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from pptx import Presentation

from cell_local.export import export_svg_to_pptx


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "source.svg"
        self.destination = self.root / "result.pptx"

    def tearDown(self):
        self.temporary.cleanup()

    def write_svg(self, body, view_box="0 0 200 100"):
        self.source.write_text(
            f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{view_box}">{body}</svg>',
            encoding="utf-8",
        )

    def test_compound_hole_is_one_native_shape_with_opposite_winding(self):
        self.write_svg(
            '<rect width="200" height="100" fill="#2463eb"/>'
            '<path fill="#ffffff" d="M0 0H200V100H0Z M20 20V80H180V20Z"/>'
        )
        result = export_svg_to_pptx(self.source, self.destination)
        deck = Presentation(self.destination)
        self.assertEqual(result["native_object_count"], 2)
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 2)
        background, ring = deck.slides[0].shapes
        self.assertEqual((background.left, background.top), (0, 0))
        self.assertEqual((background.width, background.height), (deck.slide_width, deck.slide_height))
        self.assertEqual(len(ring._element.xpath(".//a:path")), 1)
        self.assertEqual(len(ring._element.xpath(".//a:moveTo")), 2)
        self.assertEqual(len(ring._element.xpath(".//a:close")), 2)
        contours, current = [], []
        for command in ring._element.xpath(".//a:path")[0]:
            kind = command.tag.rsplit("}", 1)[-1]
            if kind in ("moveTo", "lnTo"):
                point = command[0]
                current.append((int(point.get("x")), int(point.get("y"))))
            elif kind == "close":
                contours.append(current)
                current = []
        areas = [sum(a[0] * b[1] - b[0] * a[1] for a, b in zip(c, c[1:] + c[:1])) for c in contours]
        self.assertLess(areas[0] * areas[1], 0, "The hole must retain the opposite winding")
        with ZipFile(self.destination) as package:
            self.assertFalse(any(name.startswith("ppt/media/") for name in package.namelist()))

    def test_text_is_editable_and_layer_order_is_retained(self):
        self.write_svg(
            '<rect width="200" height="100" fill="#ffffff"/>'
            '<text x="100" y="45" text-anchor="middle" font-size="12" '
            'font-family="Arial" fill="#123456">Graph labels</text>'
            '<circle cx="120" cy="60" r="5" fill="#ee4455"/>'
        )
        result = export_svg_to_pptx(self.source, self.destination)
        shapes = Presentation(self.destination).slides[0].shapes
        self.assertEqual(len(shapes), 3)
        self.assertEqual(shapes[1].text, "Graph labels")
        self.assertEqual(shapes[1].text_frame.paragraphs[0].runs[0].font.name, "Arial")
        self.assertTrue(result["warnings"])

    def test_unsupported_evenodd_does_not_overwrite_existing_file(self):
        self.write_svg('<path fill-rule="evenodd" d="M0 0H200V100H0Z M20 20H180V80H20Z"/>')
        self.destination.write_bytes(b"existing-user-file")
        with self.assertRaisesRegex(ValueError, "Even-odd"):
            export_svg_to_pptx(self.source, self.destination)
        self.assertEqual(self.destination.read_bytes(), b"existing-user-file")

    def test_raster_wrapper_is_rejected(self):
        self.write_svg('<image width="200" height="100" href="data:image/png;base64,eA=="/>')
        with self.assertRaisesRegex(ValueError, "Unsupported SVG element <image>"):
            export_svg_to_pptx(self.source, self.destination)
        self.assertFalse(self.destination.exists())


if __name__ == "__main__":
    unittest.main()
