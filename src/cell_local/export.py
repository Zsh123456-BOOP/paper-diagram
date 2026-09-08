"""Export supported vector SVGs as native, editable PowerPoint objects.

The renderer is the MIT-licensed offline backend from yrui-cmd/cell_su7.
It never embeds an SVG/bitmap and does not call an application or web service.
SVG paint order and nonzero compound contours are preserved; SVG groups are
flattened. Text stays editable, but font metrics/baselines are approximate.
Dashed strokes, gradients, masks, and unnormalized evenodd fills must be
expanded/normalized before export. See THIRD_PARTY_NOTICES.md.
"""

from __future__ import annotations

import json
import math
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from pptx import Presentation


_BACKEND = Path(__file__).parent / "_vendor" / "cell_su7"
UPSTREAM_COMMIT = "518dbc2ff929fa5500221dc2c5ffb94a7f60a657"


def _run(script: str, *arguments: object) -> str:
    result = subprocess.run(
        [sys.executable, str(_BACKEND / script), *map(str, arguments)],
        check=False, capture_output=True, text=True,
    )
    if result.returncode:
        detail = (result.stderr.strip() or result.stdout.strip())[-3000:]
        raise ValueError(f"SVG to PPTX conversion failed in {script}: {detail}")
    return result.stdout.strip()


def export_svg_to_pptx(svg_path: str | Path, pptx_path: str | Path) -> dict:
    """Create one native-object slide with the SVG viewBox ratio and no margin.

    Requires ``fonttools`` and ``python-pptx`` in the current Python runtime.
    The output is published only after package and native-object checks pass.
    Unsupported input raises ``ValueError`` without changing an existing file.
    No GUI, credentials, network connection, or Office installation is needed.
    """
    source = Path(svg_path).expanduser().resolve(strict=True)
    destination = Path(pptx_path).expanduser().resolve()
    if source == destination:
        raise ValueError("SVG input and PPTX output must be different files")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".cell-local-export-", dir=destination.parent,
    ) as temporary:
        scratch = Path(temporary)
        _run(
            "prepare_geometry_cache.py", "--input", source,
            "--output-dir", scratch / "cache", "--job-id", "local",
        )
        cache_path = scratch / "cache" / "geometry-cache.json"
        cache = json.loads(cache_path.read_text(encoding="utf-8"))
        view_box = cache["view_box"]
        if not all(math.isfinite(float(number)) for number in view_box):
            raise ValueError("SVG viewBox must contain finite numbers")
        if min(view_box[2:]) <= 0:
            raise ValueError("SVG viewBox width and height must be positive")
        candidate = scratch / "candidate.pptx"
        summary = json.loads(_run(
            "run_cell_ppt_ooxml.py", "--geometry-cache", cache_path,
            "--output-pptx", candidate,
        ))
        presentation = Presentation(candidate)
        if len(presentation.slides) != 1:
            raise ValueError("Expected exactly one generated slide")
        shapes = presentation.slides[0].shapes
        if len(shapes) != summary["native_object_count"]:
            raise ValueError("Generated native-object count does not match the saved file")
        with ZipFile(candidate) as package:
            if any(name.startswith("ppt/media/") for name in package.namelist()):
                raise ValueError("Raster/SVG media embedding is not allowed")
        if any(shape._element.tag.rsplit("}", 1)[-1] != "sp" for shape in shapes):
            raise ValueError("Output must contain only native shapes and text boxes")
        expected_ratio = float(view_box[2]) / float(view_box[3])
        actual_ratio = presentation.slide_width / presentation.slide_height
        if not math.isclose(actual_ratio, expected_ratio, rel_tol=1e-6):
            raise ValueError("Generated slide does not preserve the source canvas ratio")
        warnings = []
        if any(atom["kind"] == "text" for atom in cache["atoms"]):
            warnings.append(
                "Text remains editable; font substitution and approximate text metrics "
                "can shift labels. Inspect the final file in the target application."
            )
        if any(
            part.get("stroked")
            for atom in cache["atoms"] for part in atom.get("paintParts", [])
        ):
            warnings.append(
                "The upstream backend approximates stroke caps, joins and very thin "
                "strokes. Filled-outline paths offer the most faithful export."
            )
        summary.update({
            "output_pptx": str(destination),
            "source_svg": str(source),
            "source_view_box": view_box,
            "slide_width_emu": presentation.slide_width,
            "slide_height_emu": presentation.slide_height,
            "embedded_media_count": 0,
            "upstream_commit": UPSTREAM_COMMIT,
            "warnings": warnings,
        })
        os.replace(candidate, destination)
    return summary
